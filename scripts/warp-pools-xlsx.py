"""워프 정기카풀 텍스트 → 엑셀(누적). 사용: python3 scripts/warp-pools-xlsx.py 출력.xlsx 입력1.txt [입력2.txt ...]
같은 코드는 새 값으로 덮어쓰고, 출력 파일이 이미 있으면 그 내용에 더한다."""
import re, sys, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

CODE = re.compile(r"^[A-Z]{6,12}$"); WID = re.compile(r"^[0-9]{7}[A-Z]$", re.I)
HEAD = ["코드", "사용", "역할", "자택", "회사", "지역", "출근", "퇴근", "회원 고유번호", "등록일"]

TIME = re.compile(r"^([01]?\d|2[0-3]):[0-5]\d$")
WID2 = re.compile(r"^[0-9]{7}[A-Z0-9]$", re.I)

def blocks(text):
    recs, cur = [], []
    for l in [x.strip() for x in text.replace("\r", "").split("\n")]:
        if not l:
            if cur: recs.append(cur); cur = []
            continue
        if CODE.match(l) and cur and any(TIME.match(x) for x in cur): recs.append(cur); cur = []
        cur.append(l)
    if cur: recs.append(cur)
    return recs

def one(c):
    """[코드, 사용, 역할, 주소 1~2개, (지역), 출근, 퇴근, 고유번호, 등록일, (쪽 번호…)] → 10칸. 회사 주소·지역이 빠진 글도 받는다."""
    if not CODE.match(c[0]) or len(c) < 7: return None
    t = next((i for i in range(3, len(c)) if TIME.match(c[i])), -1)
    if t < 0:  # 출근·퇴근 시간이 빈 글: 바로 고유번호가 온다
        w = next((i for i in range(4, len(c)) if WID2.match(c[i])), -1)
        if w >= 4: c = c[:w] + ["", ""] + c[w:]; t = w
    if t < 4 or t + 2 >= len(c): return None
    # 고유번호가 빠지고 바로 등록일이 오는 글(워프 탈퇴 회원 등)도 받는다
    if re.match(r"^\d{4}-\d{2}-\d{2}$", c[t + 2]): c = c[:t + 2] + [""] + c[t + 2:]
    elif not WID2.match(c[t + 2]): return None
    mid = c[3:t]
    region = ""
    if len(mid) >= 2 and not re.search(r"\d", mid[-1]): region = mid.pop()
    elif len(mid) == 1 and not re.search(r"\d", mid[0]): return None
    if not mid: return None
    home, work = mid[0], (mid[1] if len(mid) > 1 else "")
    reg = c[t + 3] if t + 3 < len(c) and re.match(r"^\d{4}-\d{2}-\d{2}$", c[t + 3]) else ""
    return [c[0], c[1], c[2], home, work, region, c[t], c[t + 1], c[t + 2].upper(), reg]

def parse(text):
    ok, bad = [], 0
    for c in blocks(text):
        r = one(c)
        if r: ok.append(r)
        elif CODE.match(c[0]): bad += 1
    return ok, bad

out, inputs = sys.argv[1], sys.argv[2:]
rows = {}
if os.path.exists(out):
    ws = load_workbook(out)["정기카풀"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]: rows[r[0]] = [("" if v is None else str(v)) for v in r[:10]]
before = len(rows); bad = 0
for f in inputs:
    ok, b = parse(open(f, encoding="utf-8").read()); bad += b
    for r in ok: rows[r[0]] = r
wb = Workbook(); ws = wb.active; ws.title = "정기카풀"
ws.append(HEAD)
for r in sorted(rows.values(), key=lambda r: (r[9], r[0]), reverse=True): ws.append(r)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1E66F5"); c.alignment = Alignment(horizontal="center")
for i, w in enumerate([12, 6, 8, 40, 40, 10, 7, 7, 14, 12], 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
s = wb.create_sheet("요약")
vals = list(rows.values())
s.append(["항목", "수"])
s.append(["전체", len(vals)]); s.append(["운전자", sum(r[2] == "운전자" for r in vals)]); s.append(["탑승자", sum(r[2] == "탑승자" for r in vals)])
s.append([]); s.append(["지역", "수"])
from collections import Counter
for k, n in Counter(r[5] for r in vals).most_common(): s.append([k, n])
s.column_dimensions["A"].width = 22
wb.save(out)
print(f"저장: {len(rows)}건 (이번에 새로 {len(rows)-before}건, 잘려서 뺀 글 {bad}건)")
